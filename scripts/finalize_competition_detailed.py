from __future__ import annotations

import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree as ET

from competition_detailed_content import NAME, DATE, PAGES

PROJECT=Path(r'D:\table\ai_jiuye')
ROOT=PROJECT/'docs'/'deliverables'/'比赛提交材料_职航智管_20261025'
EVIDENCE=ROOT/'evidence_detailed'
TMP=PROJECT/'output'/'competition-detailed-20261025'
FINAL=ROOT/(NAME+'_作品介绍_详细版.docx')
MD=ROOT/(NAME+'_作品介绍_详细版.md')
PDF=ROOT/(NAME+'_作品介绍_详细版.pdf')


def additional_documents():
    (ROOT/'README.md').write_text(f'''# 职航智管比赛项目材料

建议先阅读 [作品介绍详细版 PDF]({PDF.name})。详细版有24章，逐项说明背景、架构、业务场景、接口状态、验证边界、开源治理和后续验收；对应 [Word]({FINAL.name}) 与 [Markdown]({MD.name}) 使用同一内容源。原15页介绍仍保留，文件名不含“详细版”。

## 详细版新增内容

- 角色边界、用户旅程、主Agent与专业Agent分工、工具授权和人工确认。
- 消息幂等、SSE事件、动作状态、失败恢复和数据模型。
- 画像计划、岗位简历、面试、学习职场、投递复盘、辅导员和入职回访。
- 全量部署与服务器演示能力对照，历史测试日期和当前站点连通检查范围。
- 许可现状、依赖来源、Issue/PR方案、版本发布、试点与长期验收计划。
- 演示脚本、异常预案、三项必交材料和源码证据定位。

## 提交前补齐

代码仓库URL与评审tag、视频URL、评审专用凭据以及报名字段。凭据通过赛事指定安全渠道提供，不能写进公开视频或公开仓库。项目许可证和依赖许可仍待实际确认，文档不代替授权决定。

## 样本与验证时间

2026-10-20历史记录包括后端63/63、前端25/25、API22/22、风险边界15/15和现有浏览器1/6；演示库是人工构造数据。2026-10-22实际DOCX上传保存成功但解析失败，材料保留此结果。2026-10-25仅核对源码、编制文档并检查服务器首页/健康接口，没有重跑全套测试或登录模型流程。

`assets_detailed/`为新稿使用的裁切截图与图；`evidence_detailed/`保存来源和历史日志。渲染中间文件不进入新压缩包。
''',encoding='utf-8')
    (ROOT/'比赛提交清单.md').write_text(f'''# 比赛提交清单

比赛截图中赛道一与赛道二要求以下三项；具体时长、格式或其他补充要求以组委会通知为准，不沿用其他赛事的限制。

## 必交材料

- [ ] 完整可运行代码的HTTPS仓库链接及评审tag/commit。
- [x] 作品介绍内容已完成：[详细版PDF]({PDF.name})，提交前最后核对链接、版本与身份信息。
- [ ] 可访问的作品演示视频URL，包含核心功能与使用流程。

## 评审内容定位

- [x] 技术创新30%：详细版04—08章；给出编排、工具、契约、状态与证据设计。
- [x] 场景落地30%：09—18章；覆盖学生、辅导员、入职后训练、部署与测试。
- [x] 开源治理20%：19—20章；现状和治理方案分开，尚未完成许可证和公开仓库。
- [x] 长期发展20%：21—22章；试点、责任与指标为计划，没有虚构团队承诺。

上面的勾选只表示文档覆盖，不能据此认定项目已完成开源或通过评审。

## 实际项目仍需处理

- [ ] 确认项目权利主体、许可证、依赖与模型分发条件。
- [ ] 修复或明确说明上传解析、风险数据采集、统计分母、详情恢复和历史E2E缺口。
- [ ] 在待提交版本复测并标注日期；不要把2026-10-20结果写成新测试。
- [ ] 确认服务器能登录并完成核心流程；当前只读首页与健康接口200不足以替代。
- [ ] 上传视频后检查公开权限，完整观看并核验与当前版本一致。
- [ ] 填写报名字段，准备专用凭据；从截图中移除个人姓名、学校、联系方式及真实简历。
- [ ] 复核仓库不含密钥、数据库、上传材料或不应分发的模型缓存。
''',encoding='utf-8')
    (ROOT/'视频提交说明.md').write_text('''# 视频与现场演示说明

比赛截图要求视频展示核心功能和使用流程，没有显示视频时长上限。建议核心演示6—8分钟是本稿编排建议，实际服从最终赛事通知。

| 顺序 | 演示内容 | 应指出的事实 |
| --- | --- | --- |
| 1 | 学生登录与工作台 | 使用的部署模式，数据是演示样例 |
| 2 | 画像、计划、岗位 | 确认状态、任务标准、匹配证据 |
| 3 | 就业助手 | 专业步骤、工具参数、来源与Provider |
| 4 | 面试与职场训练 | 单轮和场次分开；草案与实际交付分开 |
| 5 | 辅导员帮扶 | 风险事实、学生反馈与人工验收 |

模型不可用时展示有日期和Provider的已完成记录，说明是历史结果；解析失败时保留状态并展示文本输入路径，不伪造成功。服务器演示版不运行语音和视觉服务。

本地现有“职航智管.mp4”“7月23日.mp4”为候选素材，尚未在本次重新核验视频完整内容、身份信息与当前功能的一致性。先人工完整观看，再决定重录或使用。

视频公开URL：【提交前填写】。视频和PDF中不要出现评审密码、令牌、真实学生资料或可复用的企业邀请链接。
''',encoding='utf-8')
    shutil.copy2(PROJECT/'docs'/'competition-governance-proposal.md',ROOT/'开源治理与长期发展专项说明.md')


def validate():
    doc=Document(FINAL)
    all_text='\n'.join(p.text for p in doc.paragraphs)
    table_text=[['\t'.join(c.text for c in row.cells) for row in t.rows] for t in doc.tables]
    tables_blob='\n'.join('\n'.join(t) for t in table_text)
    full=all_text+'\n'+tables_blob
    assert len([p for p in doc.paragraphs if p.style.name=='Heading 1'])==len(PAGES)==24
    assert len(doc.inline_shapes)==15
    for page in PAGES:
        assert page['title'] in all_text
        for block in page['blocks']:
            if block['kind']in ['p','h','code']:assert block['text']in all_text,(page['id'],block['text'][:30])
            elif block['kind']=='figure':assert (ROOT/'assets_detailed'/(block['asset']+'.png')).is_file()
            elif block['kind']=='table':
                for row in [block['headers']]+block['rows']:
                    for value in row:assert str(value)in tables_blob
    markdown=MD.read_text(encoding='utf-8')
    for path in re.findall(r'\]\(([^)]+)\)',markdown):assert (ROOT/path).is_file(),path
    assert markdown.count('```')%2==0
    assert not re.search(r'Student123!|Counselor123!|eyJ[A-Za-z0-9_-]{40,}|sk-[A-Za-z0-9_-]{20,}',full)
    assert '\ufffd'not in full
    with zipfile.ZipFile(FINAL)as archive:
        assert archive.testzip()is None
        xml=ET.fromstring(archive.read('word/document.xml'))
        assert not xml.xpath('//*[local-name()="ins" or local-name()="del"]')
    facts=json.loads((EVIDENCE/'document_basis.json').read_text(encoding='utf-8'))
    data=json.loads((EVIDENCE/'data_summary.json').read_text(encoding='utf-8'))
    assert data['backend_tests']['tests']==63 and data['backend_tests']['failures']==0
    assert data['frontend_tests']['passed']==25
    assert data['existing_e2e']['failed']==5
    assert sum(data['current_stage_counts'].values())==85
    render=json.loads((EVIDENCE/'pdf_render_check.json').read_text(encoding='utf-8'))
    assert render['pages']==25 and render['all_headings_on_expected_pages']
    result={'prepared_on':DATE,'revision_completed_on':'2026-10-26','pages':25,'chapters':24,'figures':15,'tables':len(doc.tables),
    'nonwhitespace_text_characters':len(re.sub(r'\s','',full)),'chinese_characters':len(re.findall(r'[\u3400-\u9fff]',full)),
    'docx_markdown_share_source':True,'all_referenced_figures_present':True,
    'historical_test_dates_explicit':True,'test_suite_rerun':False,
    'repository_and_video_links_pending':True,'known_limitations_retained':True,
    'outputs':[{'file':f.name,'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()}for f in [FINAL,MD,PDF]]}
    (EVIDENCE/'delivery_check.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result


def pack():
    # Preserve directories. Earlier Compress-Archive of a file list flattened
    # folder structure, which broke local Markdown image paths after extraction.
    complete=ROOT.parent/'比赛提交材料_职航智管_20261025.zip'
    detailed=ROOT.parent/'职航智管_比赛项目文档_详细版_20261025.zip'
    excludes={'rendered','rendered-detailed'}
    with zipfile.ZipFile(complete,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6)as archive:
        for file in sorted(ROOT.rglob('*')):
            rel=file.relative_to(ROOT)
            if file.is_file()and not any(part in excludes for part in rel.parts):archive.write(file,ROOT.name+'/'+rel.as_posix())
    selected=[FINAL,MD,PDF,ROOT/'README.md',ROOT/'比赛提交清单.md',ROOT/'视频提交说明.md',ROOT/'开源治理与长期发展专项说明.md']
    selected += [p for folder in ['assets_detailed','evidence_detailed']for p in (ROOT/folder).rglob('*')if p.is_file()]
    with zipfile.ZipFile(detailed,'w',compression=zipfile.ZIP_DEFLATED,compresslevel=6)as archive:
        for file in selected:archive.write(file,'职航智管_比赛项目文档_详细版/'+file.relative_to(ROOT).as_posix())
    for file in [complete,detailed]:
        with zipfile.ZipFile(file)as archive:
            assert archive.testzip()is None
            assert any('/assets_detailed/agent.png'in name for name in archive.namelist())
    return {'complete_zip':str(complete),'detailed_zip':str(detailed),'detailed_zip_mib':round(detailed.stat().st_size/1048576,2)}


def main():
    additional_documents();result=validate();result.update(pack())
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__=='__main__':main()
