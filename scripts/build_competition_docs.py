from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

PROJECT = Path(r"D:\table\ai_jiuye")
OUTPUT = PROJECT / "docs" / "deliverables" / "比赛提交材料_职航智管_20261025"
ASSETS = OUTPUT / "assets"
DOCX = OUTPUT / "职航智管求职训练与职业成长服务系统_作品介绍.docx"
MD = OUTPUT / "职航智管求职训练与职业成长服务系统_作品介绍.md"
CHECKLIST = OUTPUT / "比赛提交清单.md"
for p in [OUTPUT, ASSETS]:
    p.mkdir(parents=True, exist_ok=True)

EVIDENCE = PROJECT / "output" / "ppt-evidence-2026-10-20"
OLD_ASSETS = PROJECT / "docs" / "deliverables" / "assets"
NAME = "职航智管求职训练与职业成长服务系统"

PURPLE = "6F35E9"
LIGHT_PURPLE = "F6F2FF"
PALE = "FBFAFF"
DARK = "211A38"
GRAY = "625C72"
GREEN = "0F766E"
ORANGE = "C2410C"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_border(cell, color="E6DFFF", size="8"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_font(run, name="等线", size=10.5, bold=False, color=DARK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for k in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn("w:" + k), name)
    return run


def add_p(doc, text="", size=10.5, bold=False, color=DARK, before=0, after=5, line=1.15, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        p.alignment = align
    set_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def add_rich_p(doc, pieces, before=0, after=5, line=1.15):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    for text, kwargs in pieces:
        set_font(p.add_run(text), **kwargs)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6 if level == 1 else 2)
    p.paragraph_format.space_after = Pt(7 if level == 1 else 4)
    p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run(text), size=20 if level == 1 else 13, bold=True, color=DARK)
    return p


def table(doc, headers, rows, widths=None, font_size=9.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        if widths:
            c.width = Cm(widths[j])
        set_cell_shading(c, PURPLE)
        set_cell_border(c, PURPLE)
        set_cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(h), size=font_size, bold=True, color="FFFFFF")
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, value in enumerate(row):
            c = cells[j]
            if widths:
                c.width = Cm(widths[j])
            set_cell_shading(c, "FFFFFF" if i % 2 == 0 else LIGHT_PURPLE)
            set_cell_border(c)
            set_cell_margins(c, top=80, bottom=80)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            set_font(p.add_run(str(value)), size=font_size, color=DARK)
    return t


def add_image(doc, path: Path, width_cm: float = 16.5, caption: str | None = None):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(5)
        set_font(cap.add_run(caption), size=8.5, color=GRAY)


def add_bullets(doc, items, size=10.2):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.12
        set_font(p.add_run(item), size=size, color=DARK)


def page_break(doc):
    doc.add_page_break()


def mask_student_image(src: Path, dest: Path):
    im = Image.open(src).convert("RGB")
    draw = ImageDraw.Draw(im)
    # Cover the demo profile/name chips while keeping the functional content.
    draw.rectangle((im.width - 300, 0, im.width, 100), fill="#F7F8F3")
    draw.rectangle((0, im.height - 120, 430, im.height), fill="#F7F8F3")
    draw.text((im.width - 260, 35), "学生 A", fill="#1F2937")
    dest.parent.mkdir(parents=True, exist_ok=True)
    im.save(dest)


def prepare_assets():
    mappings = {
        "student_dashboard.png": EVIDENCE / "screenshots" / "01-student-dashboard.png",
        "agent_assistant.png": EVIDENCE / "screenshots" / "09-agent-assistant.png",
        "counselor_dashboard.png": EVIDENCE / "screenshots" / "12-counselor-dashboard.png",
        "interview.png": EVIDENCE / "screenshots" / "07-interview-report.png",
        "jobs.png": EVIDENCE / "screenshots" / "04-job-match.png",
        "training.png": EVIDENCE / "screenshots" / "08-workplace.png",
        "test_results.png": EVIDENCE / "charts" / "C06-test-results.png",
        "tool_inventory.png": EVIDENCE / "charts" / "C08-tool-inventory.png",
        "risk_boundaries.png": EVIDENCE / "charts" / "C09-risk-boundaries.png",
        "funnel.png": EVIDENCE / "charts" / "C01-cumulative-funnel.png",
        "architecture.png": OLD_ASSETS / "technical-manual-architecture.png",
        "agent_flow.png": OLD_ASSETS / "technical-manual-agent-flow.png",
    }
    for out_name, source in mappings.items():
        if source.exists():
            shutil.copy2(source, ASSETS / out_name)
    student_src = ASSETS / "student_dashboard.png"
    if student_src.exists():
        mask_student_image(student_src, ASSETS / "student_dashboard_masked.png")


def add_header_footer(doc):
    for section in doc.sections:
        section.top_margin = Cm(1.7)
        section.bottom_margin = Cm(1.6)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)
        header = section.header
        p = header.paragraphs[0]
        p.text = "职航智管｜比赛项目介绍"
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.runs[0], size=8.5, color=PURPLE)
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "2026-10-25  |  评审材料  |  演示数据与实测边界见文中说明"
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(fp.runs[0], size=8, color=GRAY)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    add_header_footer(doc)
    styles = doc.styles
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "List Bullet"]:
        if style_name not in styles:
            styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
    styles["Normal"].font.name = "等线"
    styles["Normal"].font.size = Pt(10.5)
    # Cover
    add_p(doc, "职航智管", size=32, bold=True, color=PURPLE, before=60, after=4, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(doc, "求职训练与职业成长服务系统", size=21, bold=True, color=DARK, after=18, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(doc, "比赛作品介绍文档", size=15, bold=True, color=GRAY, after=26, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(doc, "评审维度：技术创新 30%｜场景落地 30%｜开源治理 20%｜长期发展 20%", size=10.5, color=PURPLE, after=22, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_image(doc, ASSETS / "agent_assistant.png", width_cm=16.2, caption="图 1  就业助手与 Agent 执行过程（脱敏演示界面）")
    add_p(doc, "文档版本  V1.0", size=10, color=GRAY, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(doc, "编制日期  2026-10-25", size=10, color=GRAY, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_p(doc, "提交链接：代码仓库【待填写】｜演示视频【待填写】", size=10, color=ORANGE, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    page_break(doc)

    # 1 executive summary
    heading(doc, "一 项目概览", 1)
    add_p(doc, "职航智管面向高校毕业生和就业辅导人员，把就业画像、四周计划、岗位匹配、简历优化、模拟面试、投递管理、学习成长、入职后训练和反馈回流组织成一个可追踪的业务闭环。系统的核心边界是：模型负责理解与提出建议，后端负责事实、权限、工具执行和审计。", size=11, after=8, line=1.25)
    table(doc, ["评审维度", "本项目回答的问题", "现有证据"], [
        ["技术创新 30%", "如何让 Agent 在真实业务中可控、可解释、可追溯？", "多 Agent 编排、工具目录、契约校验、持久会话、事件记录"],
        ["场景落地 30%", "是否覆盖学生、辅导员和入职后训练的实际任务？", "学生工作台、班级概览、帮扶闭环、职场交付工作台"],
        ["开源治理 20%", "代码是否能被运行、理解、复现和继续维护？", "README、Docker Compose、环境模板、测试、数据库迁移、部署文档"],
        ["长期发展 20%", "项目如何从演示版本走向可运营产品？", "版本化量表、风险边界、故障转移、数据保留策略和迭代路线"],
    ], widths=[3.2, 6.8, 7.2], font_size=9.2)
    add_p(doc, "材料边界", size=12, bold=True, color=PURPLE, before=10, after=4)
    add_bullets(doc, [
        "文中“演示数据”用于说明系统如何工作，不代表全国就业统计、真实用户规模或就业成效。",
        "当前本地测试使用 local-demo/local-rules；在线模型、OCR、语音和视觉服务是否可用以部署环境为准。",
        "代码仓库链接、评审账号和演示视频链接尚未提供，已统一留为待填写项。",
    ])
    page_break(doc)

    # 2 evaluation map
    heading(doc, "二 评审维度对照表", 1)
    add_p(doc, "根据比赛页面给出的四项评分维度，本作品将功能证据、技术实现和后续治理放在同一条证据链中。", size=11, after=8, line=1.25)
    table(doc, ["维度", "评审关注点", "项目回应", "建议展示材料"], [
        ["技术创新 30%", "是否解决明确问题，是否有技术优势", "证据驱动的多 Agent 协作；后端工具授权；流式持久会话；面试和职场训练分域", "架构图、Agent 执行截图、工具目录、代码仓库"],
        ["场景落地 30%", "是否真实、可运行、可测试", "学生执行、辅导员帮扶、入职后交付三条场景链；支持浏览器操作和结构化结果", "学生/辅导员截图、演示视频、测试记录"],
        ["开源治理 20%", "协议、依赖、文档、Issue/PR 是否清晰", "已有 README、环境变量模板、Docker Compose、迁移脚本和自动化测试；公共仓库待填写", "README、部署文档、测试结果、仓库链接"],
        ["长期发展 20%", "持续维护、治理和组织支持", "风险规则、量表版本、数据保留、故障转移和分阶段路线；上线前仍需补充运营指标", "路线图、风险边界、版本计划"],
    ], widths=[2.7, 4.6, 6.3, 3.6], font_size=8.6)
    add_image(doc, ASSETS / "test_results.png", width_cm=16.0, caption="图 2  本次验证结果：后端 63/63、前端组件 25/25、API 流程 22/22、风险边界 15/15；现有浏览器回归仍有 5 项待处理")
    page_break(doc)

    # 3 technical innovation
    heading(doc, "三 技术创新：把 Agent 放进可控的业务边界", 1)
    add_p(doc, "本项目的创新重点不是把模型回答做得更像聊天，而是把模型放进一个由业务事实、权限、工具和审计共同约束的运行环境。", size=11, after=8, line=1.25)
    add_image(doc, ASSETS / "architecture.png", width_cm=16.0, caption="图 3  系统架构：学生/辅导员工作区、Agent、业务服务和数据层")
    table(doc, ["创新点", "实现方式", "评审可验证证据"], [
        ["证据驱动", "画像、岗位、任务、申请和面试结果作为业务事实；回答携带来源片段", "就业助手回答与引用；岗位匹配五维证据；每周复盘读取事件"],
        ["可控编排", "主 Agent 规划步骤，专业 Agent 处理单一职责，不能直接改数据库", "Agent 运行详情、工具目录和专业 Agent 步骤"],
        ["授权工具", "角色、资源归属、参数契约和风险等级由后端重新校验", "20 个只读、9 个普通写、3 个高风险拒绝入口的目录盘点"],
        ["持久可追溯", "会话、消息、运行状态、工具调用和动作提案写入服务端", "刷新后消息保持；工具输入输出、耗时和失败状态可查看"],
    ], widths=[3.0, 8.0, 6.2], font_size=8.8)
    page_break(doc)

    # 4 agent implementation
    heading(doc, "四 技术创新：Agent、工具和风险等级", 1)
    add_image(doc, ASSETS / "agent_flow.png", width_cm=16.0, caption="图 4  Agent 运行流程：问题理解、专业 Agent、授权工具、业务状态和持久事件")
    add_p(doc, "系统允许模型提出工具候选，但不允许模型创建工具、扩大权限或绕过后端规则。普通写操作可以进入动作提案，高风险删除、批量修改和代替学生向外部平台提交等动作保持二次确认或拒绝。", size=10.8, after=7, line=1.2)
    add_image(doc, ASSETS / "tool_inventory.png", width_cm=14.5, caption="图 5  注册工具按风险分层：只读 20、普通写 9、高风险拒绝入口 3")
    add_bullets(doc, [
        "结构化契约：Agent 计划、工具参数、工具输出和回答均有 Pydantic 契约。",
        "Provider 故障转移：在线模型不可用时可使用本地规则，但不会把失败伪装成在线模型成功。",
        "安全边界：摄像头表达分析只作训练参考，不用于心理诊断、录用判断、性别判断或外貌评价。",
    ])
    page_break(doc)

    # 5 scenario student
    heading(doc, "五 场景落地：学生从目标到行动", 1)
    add_image(doc, ASSETS / "student_dashboard_masked.png", width_cm=16.2, caption="图 6  学生工作台：目标岗位、任务、投递漏斗和帮扶反馈集中展示（演示数据）")
    table(doc, ["用户任务", "页面链路", "输出结果"], [
        ["建立可核对的画像", "材料/文本 → 能力证据 → 待确认画像 → 正式档案", "优势、差距、约束和来源片段"],
        ["执行四周计划", "目标岗位 → 计划 → 周任务 → 截止日期 → 验收标准", "可执行任务，而不是泛化建议"],
        ["完成岗位准备", "岗位库 → 五维匹配 → 简历建议 → 模拟面试", "匹配证据、补齐动作和训练反馈"],
        ["形成复盘闭环", "投递事件 → 每周复盘 → 下周三项行动", "下一轮行动有真实记录依据"],
    ], widths=[4.0, 7.2, 6.0], font_size=9.0)
    add_p(doc, "演示样本说明：本地演示库包含 24 名构造学生、85 条申请记录和 96 项计划任务；这些数字不代表实际用户或就业统计。", size=9.2, color=GRAY, after=0)
    page_break(doc)

    # 6 counselor
    heading(doc, "六 场景落地：辅导员从风险到帮扶", 1)
    add_image(doc, ASSETS / "counselor_dashboard.png", width_cm=16.2, caption="图 7  辅导员班级概览：学生数、申请记录和风险事实（演示数据）")
    add_p(doc, "辅导员端将学生状态从“人工收集”转为“事实可查、风险可解释、任务可验收”。风险规则只提供线索，辅导员仍需查看触发事实并人工确认，不自动给学生排名、处分或限制资源。", size=10.8, after=7, line=1.2)
    table(doc, ["环节", "系统动作", "人工责任"], [
        ["事实读取", "读取任务、投递、面试和反馈记录", "确认记录是否完整、是否需要补充"],
        ["风险提示", "根据无投递、任务逾期等规则生成风险事实", "判断是否需要帮扶，不把规则当结论"],
        ["任务闭环", "创建帮扶任务，学生提交反馈，辅导员确认完成", "核验反馈和结果证据，必要时继续跟进"],
        ["质量回流", "入职回访、企业评价和课程改进数据分开保存", "按权限查看并解释数据使用范围"],
    ], widths=[3.0, 8.2, 6.0], font_size=9.0)
    add_p(doc, "本次补充浏览器检查完成了“学生提交帮扶反馈 → 辅导员确认完成”，移动端视口也检查了无横向溢出。", size=9.2, color=GRAY, after=0)
    page_break(doc)

    # 7 training
    heading(doc, "七 场景落地：面试训练与入职后工作训练", 1)
    add_image(doc, ASSETS / "interview.png", width_cm=8.1, caption="图 8  模拟面试场次报告：6 道核心题 + 2 次追问")
    add_image(doc, ASSETS / "training.png", width_cm=8.1, caption="图 9  职场工作台：任务简报、验收标准、交付物与复盘")
    table(doc, ["模块", "解决的问题", "边界"], [
        ["模拟面试", "岗位上下文、动态追问、流式回答和四维评分", "摄像头可选；评分仅供表达训练，不代表录用判断"],
        ["职场训练", "任务澄清、行动计划、交付物提交和结构化评审", "AI 只能生成工作草案，学生核对后主动提交"],
        ["学习成长", "能力差距、课程、训练任务和考试结果", "完成状态必须由实际学习记录或证据支持"],
    ], widths=[3.0, 8.0, 6.2], font_size=9.0)
    add_p(doc, "产品区分了入职前模拟面试和入职后职场训练，避免把面试回答、工作交付物和就业结果混为一谈。", size=9.2, color=GRAY, after=0)
    page_break(doc)

    # 8 engineering
    heading(doc, "八 可运行性与工程交付", 1)
    add_p(doc, "项目提供 Windows 本地启动脚本和 Docker Compose 配置，后端使用 FastAPI，前端使用 React/TypeScript，数据层支持 SQLite 本地演示或 PostgreSQL 部署。异步 AI 任务由 Worker 和 Redis 承载，OCR、向量、语音和视觉能力按服务状态启用。", size=10.8, after=8, line=1.2)
    table(doc, ["工程项", "当前实现", "提交前说明"], [
        ["环境配置", ".env.example、local launcher、Docker Compose", "公共部署前必须替换密钥和演示密码"],
        ["数据与迁移", "SQLAlchemy 模型、Alembic 迁移、演示种子", "迁移状态需在目标环境复验"],
        ["测试验证", "后端 63/63、前端组件 25/25、API 22/22、规则边界 15/15", "现有浏览器回归仍有 5 项待修订"],
        ["可观测性", "审计日志、Agent 用量、运行状态、错误与耗时", "生产环境需接入持续监控与告警"],
        ["隐私控制", "JWT/RBAC、资源归属、日志脱敏、摄像头原始帧不落库", "上线前补充数据保留期限与主体授权流程"],
    ], widths=[3.0, 8.3, 5.9], font_size=8.8)
    add_image(doc, ASSETS / "architecture.png", width_cm=15.0, caption="图 10  工程架构与部署边界")
    page_break(doc)

    # 9 open source governance
    heading(doc, "九 开源治理：让代码可运行、可理解、可继续维护", 1)
    add_p(doc, "比赛要求提交代码仓库链接。当前工作区没有配置公开 Git remote，因此材料不伪造仓库地址；提交前应将仓库链接补入清单，并确认仓库内容与演示版本一致。", size=10.8, after=8, line=1.2)
    table(doc, ["治理对象", "当前已有", "发布前动作"], [
        ["运行文档", "README、Windows 启动脚本、Docker Compose、部署说明", "在仓库首页放置一条可复现启动路径"],
        ["依赖与配置", "backend/frontend package 配置、.env.example", "禁止提交密钥、数据库文件、个人简历和演示密码"],
        ["质量门禁", "pytest、Vitest、Playwright、构建脚本", "把核心测试接入 CI，记录失败原因而非只展示通过数"],
        ["Issue/PR", "当前未提供公开仓库链接", "建立 bug、feature、security 模板，要求复现步骤和影响范围"],
        ["版本与变更", "源码内有本地版本与迁移记录", "为评审版本打 tag，发布 CHANGELOG 和已知限制"],
    ], widths=[3.3, 7.8, 6.1], font_size=8.8)
    add_image(doc, ASSETS / "test_results.png", width_cm=15.3, caption="图 11  质量验证记录；浏览器回归未通过项已保留在项目说明中")
    page_break(doc)

    # 10 long-term development
    heading(doc, "十 长期发展：从演示闭环到可运营服务", 1)
    table(doc, ["阶段", "目标", "可验收结果"], [
        ["阶段一｜提交前", "完成公开仓库、评审链接、脱敏素材和可复现演示", "评委可打开仓库、PDF 和视频；部署说明能在干净环境启动"],
        ["阶段二｜试点", "与高校就业部门共同校准风险规则、面试量表和权限流程", "形成授权数据字典、人工复核记录和版本化量表"],
        ["阶段三｜运营", "完善岗位来源适配、任务队列、监控和数据保留", "有来源新鲜度、失败重试、告警、成本与审计指标"],
        ["阶段四｜扩展", "支持更多学校和角色，沉淀课程改进与就业质量反馈", "多租户隔离、租户级权限、跨班级汇总和可解释导出"],
    ], widths=[3.3, 8.0, 5.9], font_size=9.0)
    add_p(doc, "长期发展不以扩大 Agent 自动执行权限为目标，而以更可靠的事实、清晰的授权、更低的重复工作和可审计的服务质量为目标。", size=11, bold=True, color=PURPLE, before=12, after=8, line=1.2)
    add_image(doc, ASSETS / "funnel.png", width_cm=15.2, caption="图 12  演示申请阶段数据：仅用于说明流程分析，不代表真实就业转化")
    page_break(doc)

    # 11 risks
    heading(doc, "十一 风险、边界与改进", 1)
    table(doc, ["风险", "当前控制", "仍需补强"], [
        ["个人材料泄露", "JWT/RBAC、资源归属校验、日志脱敏、摄像头原始帧不保存", "上线前完成主体授权、数据保留和删除流程"],
        ["模型编造经历", "引用原文、草稿确认、结构化契约、未匹配来源时提示", "扩充证据评测集和人工抽检机制"],
        ["Agent 越权操作", "工具白名单、角色/资源/参数/风险校验、二次确认", "公开仓库后增加安全审查和威胁模型"],
        ["外部岗位过期", "保存来源 URL、时间和可信等级，提示原页面核验", "增加适配器监控、过期处理和失败告警"],
        ["服务资源消耗", "Worker、超时、重试、任务状态、死信记录", "增加并发压测、成本预算和容量监控"],
    ], widths=[3.1, 8.2, 5.9], font_size=8.8)
    add_image(doc, ASSETS / "risk_boundaries.png", width_cm=15.6, caption="图 13  五条风险规则的阈值边界验证：15/15 通过")
    page_break(doc)

    # 12 demo and submission
    heading(doc, "十二 演示路径与提交清单", 1)
    add_p(doc, "推荐现场演示顺序：登录 → 学生工作台 → 就业画像/四周计划 → 岗位匹配 → 就业助手执行详情 → 模拟面试 → 职场训练 → 辅导员概览 → 帮扶验收。每段只展示一个可验证结果，避免等待外部服务。", size=10.8, after=8, line=1.2)
    table(doc, ["提交项", "当前状态", "最终填写/检查"], [
        ["代码仓库链接", "待填写；当前项目未配置公开 remote", "补充 HTTPS 仓库 URL，并确认可访问、可运行、无密钥"],
        ["作品介绍文档", "本 PDF / DOCX", "上传 PDF，控制文件大小，确认图片和文字可读"],
        ["演示视频链接", "待填写；本地存在职航智管.mp4 和 7月23日.mp4", "上传指定视频平台，补充可访问链接和时长"],
        ["账号与权限", "本地演示账号仅供开发环境", "在评审环境创建专用账号，禁止把密码写入公开视频或公开仓库"],
        ["脱敏检查", "现有截图已使用演示数据，仍需逐页复核", "检查姓名、学校、教师、学号、手机号、邮箱和真实简历原文"],
    ], widths=[4.0, 7.2, 6.0], font_size=9.0)
    add_p(doc, "评审材料提交地址：代码【待填写】｜介绍 PDF【当前文档】｜演示视频【待填写】", size=11, bold=True, color=PURPLE, before=12, after=8, line=1.15)
    add_image(doc, ASSETS / "agent_assistant.png", width_cm=16.0, caption="图 14  现场讲解重点：回答、依据和工具执行过程在同一界面")
    page_break(doc)

    # appendix
    heading(doc, "附录 事实口径与证据索引", 1)
    add_bullets(doc, [
        "测试记录来源：output/ppt-evidence-2026-10-20；后端 63/63、前端组件 25/25、API 22/22、风险规则边界 15/15。",
        "演示库口径：24 名人工构造学生、6 个演示岗位、85 条申请记录、96 项计划任务、9 条风险事件。",
        "匹配与面试分数是 local-demo/local-rules 演示输出，不是算法准确率、录用概率或真实能力提升。",
        "现有浏览器回归仍有 5 项待处理：多记录选择器、旧断言文案和移动端职场流程超时；未在本文中隐藏。",
        "本地运行环境、代码仓库和视频链接不是公开提交地址；最终提交前由项目方替换待填写项。",
    ], size=10.2)
    table(doc, ["附件", "用途"], [
        ["README.md / docs/workplace-workbench.md", "运行方式、工作台边界和工程说明"],
        ["docs/deliverables/assets/", "架构图、Agent 流程图和产品截图"],
        ["output/ppt-evidence-2026-10-20/", "测试日志、JSON 数据和脱敏截图"],
        ["职航智管.mp4 / 7月23日.mp4", "本地演示视频，上传后补充公开链接"],
    ], widths=[8.0, 9.2], font_size=9.2)
    return doc


def write_markdown():
    text = f"""# {NAME}｜比赛项目介绍\n\n**评审维度：**技术创新 30%｜场景落地 30%｜开源治理 20%｜长期发展 20%  \n**版本：**V1.0  \n**编制日期：**2026-10-25\n\n## 1. 项目定位\n\n职航智管面向高校毕业生和就业辅导人员，把就业画像、四周计划、岗位匹配、简历优化、模拟面试、投递管理、学习成长、入职后训练和反馈回流组织成一个可追踪的业务闭环。模型负责理解与提出建议，后端负责事实、权限、工具执行和审计。\n\n## 2. 评审维度对照\n\n| 维度 | 项目回应 | 证据 |\n| --- | --- | --- |\n| 技术创新 30% | 证据驱动的多 Agent 协作、授权工具、持久会话、结构化契约 | 架构图、Agent 流程、工具目录、执行截图 |\n| 场景落地 30% | 学生执行、辅导员帮扶、入职后工作训练 | 产品截图、实际浏览器检查、视频 |\n| 开源治理 20% | README、Docker Compose、环境模板、迁移和测试 | 代码仓库、部署文档、测试结果 |\n| 长期发展 20% | 版本化量表、风险边界、故障转移和分阶段路线 | 风险表、路线图、已知限制 |\n\n## 3. 技术创新\n\n- **证据优先：**画像、岗位、任务、申请和面试结果作为业务事实，回答携带来源片段。\n- **可控编排：**主 Agent 规划步骤，专业 Agent 处理单一职责，不能直接改数据库。\n- **授权工具：**后端重新检查角色、资源归属、参数契约和风险等级。\n- **持久可追溯：**会话、消息、运行状态、工具调用和动作提案写入服务端。\n- **训练边界：**模拟面试与职场训练分域，摄像头分析仅用于表达训练参考。\n\n## 4. 场景落地\n\n### 学生\n\n学生可以从就业画像和简历材料开始，确认能力证据和目标岗位，生成四周计划，在岗位匹配、简历优化、模拟面试、投递看板和每周复盘之间持续推进。\n\n### 辅导员\n\n辅导员可以查看班级概览、风险事实和学生进展，建立帮扶任务，接收学生反馈并确认完成。风险规则提供线索，不自动给学生排名或限制资源。\n\n### 入职后\n\n职场工作台从工作简报读取岗位、阶段、目标、协作对象、截止时间和验收标准，学生形成真实交付物，提交结构化评审和复盘。\n\n## 5. 开源治理\n\n当前工作区没有配置公开 Git remote，因此代码仓库链接暂留：`【提交前填写仓库 HTTPS URL】`。发布前应完成：\n\n1. README、部署文档、环境变量模板和快速启动路径同步到仓库首页。\n2. 删除密钥、数据库文件、个人简历、演示密码和本地输出。\n3. 建立 Issue、PR 和安全问题模板，要求复现步骤、影响范围和验证结果。\n4. 为评审版本打 tag，附 CHANGELOG 和已知限制。\n5. 接入 CI，至少运行后端测试、前端测试、构建和关键 E2E。\n\n## 6. 长期发展\n\n- **阶段一：**完成公开仓库、评审链接、脱敏素材和可复现演示。\n- **阶段二：**与高校就业部门共同校准风险规则、面试量表和权限流程。\n- **阶段三：**完善岗位来源适配、任务队列、监控、成本和数据保留。\n- **阶段四：**支持多学校和多角色，形成课程改进与就业质量反馈。\n\n## 7. 测试与数据边界\n\n本次记录显示后端 63/63、前端组件 25/25、API 流程 22/22、风险规则边界 15/15；现有浏览器回归仍有 5 项待处理。演示库包含 24 名人工构造学生、6 个演示岗位、85 条申请记录、96 项计划任务和 9 条风险事件。这些数字不代表真实用户、真实投递或就业成效。\n\n匹配分和面试分是本地规则演示输出，不是算法准确率、录用概率或能力提升率。在线模型、OCR、语音和视觉服务的可用性取决于部署环境。\n\n## 8. 提交清单\n\n- 代码仓库：`【提交前填写】`\n- 作品介绍 PDF：当前文档导出 PDF\n- 演示视频：`【提交前填写视频链接】`\n- 评审账号：`【提交前填写专用账号】`\n- 评审密码：`【提交前填写专用密码】`\n- 脱敏复核：姓名、学校、教师、学号、手机号、邮箱、真实简历和账号密码均不得出现在公开材料中。\n"""
    MD.write_text(text, encoding="utf-8")


def write_checklist():
    CHECKLIST.write_text(f"""# 比赛提交清单｜{NAME}\n\n## 必交材料\n\n- [ ] 代码仓库链接：`【提交前填写仓库 HTTPS URL】`\n- [ ] 作品介绍文档 PDF：`{DOCX.with_suffix('.pdf').name}`\n- [ ] 演示视频链接：`【提交前填写视频链接】`\n\n## 评审维度覆盖\n\n- [x] 技术创新 30%：多 Agent 编排、工具授权、证据链、持久会话。\n- [x] 场景落地 30%：学生、辅导员、模拟面试、入职后职场训练。\n- [x] 开源治理 20%：README、部署、环境变量、迁移、测试、风险边界。\n- [x] 长期发展 20%：版本化量表、监控、数据保留、多租户路线。\n\n## 提交前必须替换\n\n- [ ] 仓库 URL 和版本 tag。\n- [ ] 视频公开链接、时长和访问权限。\n- [ ] 专用评审账号与密码。\n- [ ] 团队、学校、指导教师等报名字段。\n- [ ] 所有截图和视频中的姓名、学校、账号、联系方式、真实简历原文。\n\n## 已知限制\n\n当前工作区没有公开 Git remote；现有浏览器回归仍有 5 项待处理。项目材料保留该事实，没有将本地路径写成公开链接，也没有把演示数据写成真实成效。\n""", encoding="utf-8")


def main():
    prepare_assets()
    doc = build_docx()
    doc.save(DOCX)
    write_markdown()
    write_checklist()
    print(DOCX)
    print(MD)
    print(CHECKLIST)


if __name__ == "__main__":
    main()
