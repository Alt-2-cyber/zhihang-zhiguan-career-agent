"""Build a figure-rich competition report without replacing previous editions."""
from __future__ import annotations

import hashlib
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from competition_illustrated_assets import OUT, ASSETS, EVIDENCE, TMP, PROJECT
from competition_illustrated_content import NAME, REVISION, PAGES, used

DOCX=OUT/(NAME+'_比赛项目文档_图文增强版.docx')
PDF=DOCX.with_suffix('.pdf')
MD=DOCX.with_suffix('.md')
FIGURES=[]


def font(run,size=10.7,bold=False,color='161A22',name='等线'):
    run.font.name=name;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
    props=run._element.get_or_add_rPr();rf=props.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');props.insert(0,rf)
    for key in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+key),name)
    return run


def paragraph(doc,text='',size=10.7,bold=False,color='161A22',after=4.0,line=16.6,style=None,keep=False):
    p=doc.add_paragraph(style=style);pf=p.paragraph_format
    pf.space_before=Pt(0);pf.space_after=Pt(after);pf.line_spacing=Pt(line);pf.keep_with_next=keep;pf.widow_control=False
    for i,part in enumerate(text.splitlines()or['']):
        if i:p.add_run().add_break()
        font(p.add_run(part),size,bold,color)
    return p


def table(doc,block):
    heads,rows=block['headers'],block['rows'];widths=block.get('widths')or[1]*len(heads)
    widths=[w*17.4/sum(widths)for w in widths]
    tab=doc.add_table(rows=1,cols=len(heads));tab.autofit=False;tab.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,width in zip(tab.columns,widths):col.width=Cm(width)
    for i,row in enumerate([heads]+rows):
        cells=tab.rows[0].cells if i==0 else tab.add_row().cells
        trpr=tab.rows[i]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,val in enumerate(row):
            cell=cells[j];cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'EBE7F5'if i==0 else 'FFFFFF'if i%2 else'F6F7F9');props.append(fill)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge)
                for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+key),value)
                borders.append(e)
            props.append(borders);margin=OxmlElement('w:tcMar')
            for edge,amount in [('top',62),('bottom',62),('left',90),('right',90)]:
                node=OxmlElement('w:'+edge);node.set(qn('w:w'),str(amount));node.set(qn('w:type'),'dxa');margin.append(node)
            props.append(margin)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(13.1)
            p.paragraph_format.keep_with_next=i==0
            font(p.add_run(str(val)),9.1,i==0)
    paragraph(doc,'',size=1,line=2,after=2)


def figure(doc,block,page_id):
    path=ASSETS/(block['asset']+'.png')
    with Image.open(path)as im:w,h=im.size
    width=min(block.get('width',17.1),block.get('height',6.7)*w/h)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(1.5);p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Cm(width))
    num=len(FIGURES)+1;caption=f'图 {num:02}  '+block['caption']
    prop=p._p.xpath('.//wp:docPr')[0];prop.set('descr',caption);prop.set('name',block['asset'])
    cap=paragraph(doc,caption,size=8.5,color='545B67',after=6.5,line=12.1)
    cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    FIGURES.append({'number':num,'asset':block['asset'],'caption':block['caption'],'page_id':page_id,'width_cm':round(width,2),'height_cm':round(width*h/w,2)})


def field(p,instruction):
    for kind in ['begin','instruction','separate','cache','end']:
        if kind=='instruction':
            e=OxmlElement('w:instrText');e.set(qn('xml:space'),'preserve');e.text=' '+instruction+' '
        elif kind=='cache':e=OxmlElement('w:t');e.text='1'
        else:e=OxmlElement('w:fldChar');e.set(qn('w:fldCharType'),kind)
        r=OxmlElement('w:r');r.append(e);p._p.append(r)


def build():
    FIGURES.clear();doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8);sec.top_margin=Cm(1.55);sec.bottom_margin=Cm(1.55)
    sec.header_distance=Cm(.65);sec.footer_distance=Cm(.65)
    doc.core_properties.title=NAME+' 比赛项目文档图文增强版';doc.core_properties.author='';doc.core_properties.last_modified_by=''
    for name in ['Normal','Title','Heading 1','Heading 2']:
        if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        st=doc.styles[name];st.font.name='等线';st.font.color.rgb=RGBColor(0,0,0)
        for border in st.element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(10.7)
    header=sec.header.paragraphs[0];header.paragraph_format.space_after=Pt(0)
    font(header.add_run('职航智管  比赛项目文档图文增强版'),8.2,color='535963')
    footer=sec.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(Cm(17.2),WD_TAB_ALIGNMENT.RIGHT)
    font(footer.add_run(REVISION+'  图文修订  历史验证日期见正文\t'),8.0,color='535963');field(footer,'PAGE')
    pre=paragraph(doc,'技术创新 30%    场景落地 30%    开源治理 20%    长期发展 20%',size=9.4,color='6043B5',after=22,line=15)
    pre.paragraph_format.space_before=Pt(26)
    title=paragraph(doc,NAME,size=23,bold=True,after=10,line=32,style='Title')
    paragraph(doc,'比赛项目文档图文增强版',size=18,bold=True,after=14,line=24)
    paragraph(doc,'V1.0  ·  24章业务与技术说明  ·  16个专题图解  ·  60张图',size=10.7,color='535963',after=12)
    figure(doc,{'asset':'s13_agent','caption':'就业助手真实回放：回答、引用与工具执行同屏，local-rules及不可用服务状态保留','width':17.2,'height':9.0},'cover')
    paragraph(doc,'本版在详细文档上增加架构、时序、状态机、部署、风险与治理图，以及真实页面的局部读图说明。截图来自2026-10-20和2026-10-22的脱敏演示，新增流程图依据源码绘制，路线与治理图明确标为建议。',after=8)
    paragraph(doc,'修订日期 '+REVISION+'；历史测试日期和未完成项继续保留。代码仓库URL、视频URL与评审版本由项目方提交前确认。',size=9.8,color='535963')
    # Reading map and the second, unique cover-related overview graphic.
    nav=paragraph(doc,'阅读导航与图示约定',size=17,bold=True,line=24,after=9,style='Heading 1')
    nav.paragraph_format.page_break_before=True
    paragraph(doc,'建议先看总览与技术设计，再看配有真实页面的场景，最后查看部署、验证、治理和提交要求。同一主题的图解页紧跟正文，便于对照实现方式与实际表现。')
    figure(doc,{'asset':'d01_journey','caption':'本项目的证据和反馈闭环：图示设计不等于真实就业成效','width':17.2,'height':4.2},'navigation')
    table(doc,{'headers':['阅读段落','主要内容','图像类型'],'rows':[['01—04','背景、角色和总体架构','需求图、角色边界、架构、真实工作台'],['05—08','Agent、工具、会话和数据','时序、授权、提案与任务状态、实体关系'],['09—15','求职训练、帮扶和回访','页面读图、评分、看板、反馈与验收'],['16—18','部署、验证和已知问题','服务拓扑、历史图表、真实失败路径'],['19—24','许可、协作、路线和提交','明确标注为方案的流程与验收图']],'widths':[2.8,6.6,8.0]})
    paragraph(doc,'截图：保留产品字段、数字、失败或未就绪状态，仅裁掉无关区域。数据图：保留原日期、分母和人工构造样本标记。流程图：表示实现逻辑或明确标注的后续方案，不增加不存在的控件或能力。',after=9)
    for page in PAGES:
        p=paragraph(doc,page['title'],size=16.8,bold=True,after=9,line=24,style='Heading 1')
        p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True
        if page['kind']=='visual_detail':paragraph(doc,'图解页  ·  对照输入、操作、结果和边界',size=9,color='6043B5',after=7,line=13)
        for b in page['blocks']:
            if b['kind']=='p':paragraph(doc,b['text'])
            elif b['kind']=='h':paragraph(doc,b['text'],size=11.7,bold=True,after=3.5,line=18.2,style='Heading 2',keep=True)
            elif b['kind']=='table':table(doc,b)
            elif b['kind']=='figure':figure(doc,b,page['id'])
            elif b['kind']=='code':
                q=paragraph(doc,b['text'],size=8.8,color='424650',after=6,line=13.1)
                q.paragraph_format.keep_together=True
                for r in q.runs:r.font.name='Consolas'
        q=paragraph(doc,'实现或证据依据：'+'；'.join(page['sources']),size=7.9,color='626975',after=0,line=10.6)
        q.paragraph_format.space_before=Pt(3)
    doc.save(DOCX)
    assert len(FIGURES)==60 and len(set(f['asset']for f in FIGURES))==60
    (EVIDENCE/'figure_index.json').write_text(json.dumps(FIGURES,ensure_ascii=False,indent=2),encoding='utf-8')


def markdown():
    idx={f['asset']:f for f in FIGURES}
    def picture(asset):
        f=idx[asset];return f"![图{f['number']:02} {f['caption']}](assets/{asset}.png)"
    lines=['# '+NAME+' 比赛项目文档图文增强版','',f'修订 {REVISION}  V1.0  24章  16个专题图解  60张图','',picture('s13_agent'),'','## 阅读导航与图示约定','',picture('d01_journey'),'','本版保留详细正文，并为每个主要章节配置图解；历史测试和构造样本均标明日期与边界。','']
    for page in PAGES:
        lines+=['## '+page['title'],'']
        for b in page['blocks']:
            if b['kind']=='p':lines+=[b['text'],'']
            elif b['kind']=='h':lines+=['### '+b['text'],'']
            elif b['kind']=='figure':lines+=[picture(b['asset']),'']
            elif b['kind']=='code':lines+=['```text',b['text'],'```','']
            elif b['kind']=='table':
                esc=lambda x:str(x).replace('|','\\|').replace('\n','<br>')
                lines+=['| '+' | '.join(map(esc,b['headers']))+' |','| '+' | '.join(['---']*len(b['headers']))+' |']
                lines+=['| '+' | '.join(map(esc,row))+' |'for row in b['rows']];lines+=['']
        lines+=['实现或证据依据：'+'；'.join('`'+s+'`'for s in page['sources']),'']
    MD.write_text('\n'.join(lines),encoding='utf-8')


def evidence():
    prior=PROJECT/'docs'/'deliverables'/'比赛提交材料_职航智管_20261025'/'evidence_detailed'
    for name in ['data_summary.json','backend-tests-current.xml','backend-tests-current.log','frontend-tests.log','api_performance.json','architecture_inventory.json','browser-regression-summary.json']:
        shutil.copy2(prior/name,EVIDENCE/name)
    files=sorted({s for p in PAGES for s in p['sources']if(PROJECT/s).is_file()})
    records=[{'path':s,'sha256':hashlib.sha256((PROJECT/s).read_bytes()).hexdigest()}for s in files]
    d=Document(DOCX);text='\n'.join(p.text for p in d.paragraphs)
    for tab in d.tables:text+='\n'+'\n'.join(c.text for r in tab.rows for c in r.cells)
    stats={'revision':REVISION,'content_source':'scripts/competition_illustrated_content.py','chapter_count':24,'detail_spreads':16,'planned_body_spreads':40,'figures':60,'unique_figures':60,'diagrams':32,'screenshots':19,'data_charts':9,'nonwhitespace_characters':len(re.sub(r'\s','',text)),'chinese_characters':len(re.findall(r'[\u3400-\u9fff]',text)),'historical_test_date':'2026-10-20','historical_upload_date':'2026-10-22','new_test_suite_run':False,'new_screenshots_captured':False,'source_files':records}
    (EVIDENCE/'document_basis.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf-8')
    OUT.joinpath('README.md').write_text(f'''# 职航智管比赛文档图文增强版

建议阅读 [{PDF.name}]({PDF.name})，需要修改时使用对应Word或Markdown。原25页详细版仍保留在相邻目录。

本版保留24章详细说明，增加16个专题图解页，共60张不同的图：32张架构/流程/状态图、19张真实界面或注释图、9张历史数据图。每图有编号、说明和证据来源。Word与Markdown使用同一内容源。

界面来自2026-10-20及2026-10-22的历史脱敏演示；本次没有编造新业务数据或实测结果。匹配/面试分数是本地规则演示，路线、贡献与许可图明确作为方案，测试日期和未完成项均保留。

提交前填写代码仓库、版本tag与视频URL，凭据通过赛事指定渠道提供。`assets/`是可复用原图，`evidence/`保存图注索引、来源和历史验证。流程图不是产品界面，注释截图中的解释文字也不作为新控件。
''',encoding='utf-8')
    OUT.joinpath('提交前核对.md').write_text('''# 提交前核对

- [ ] 填写完整可运行代码仓库HTTPS地址和评审tag。
- [x] 作品介绍PDF、可编辑Word与Markdown已按同一内容源整理。
- [ ] 上传演示视频并复核内容与评审版本一致。
- [ ] 检查真实身份、凭据、简历、邀请链接是否已脱敏。
- [ ] 明确历史演示、源码实现、本次文档核对和后续计划的区别。
- [ ] 核查许可证与依赖分发条件，处理或说明已知缺口后对提交版本复测。

比赛截图没有给出文档页数或视频时长上限；最终以组委会通知为准。
''',encoding='utf-8')
    print(json.dumps({k:v for k,v in stats.items()if k!='source_files'},ensure_ascii=False,indent=2))


def main():
    build();markdown();evidence()
    print(DOCX)


if __name__=='__main__':main()
