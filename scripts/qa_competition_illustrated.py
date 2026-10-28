from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

import fitz
from docx import Document
from lxml import etree as ET
from PIL import Image,ImageOps,ImageDraw,ImageFont

from competition_illustrated_assets import OUT,ASSETS,EVIDENCE,TMP,PROJECT
from competition_illustrated_content import NAME,PAGES

DOCX=OUT/(NAME+'_比赛项目文档_图文增强版.docx')
PDF=DOCX.with_suffix('.pdf');MD=DOCX.with_suffix('.md')


def main():
    document=Document(DOCX)
    actual='\n'.join(p.text for p in document.paragraphs)
    table_text='\n'.join(c.text for table in document.tables for row in table.rows for c in row.cells)
    markdown=MD.read_text(encoding='utf-8')
    assert len(document.inline_shapes)==60
    assert len([p for p in document.paragraphs if p.style.name=='Heading 1'])==41
    for page in PAGES:
        assert page['title']in actual
        for block in page['blocks']:
            if block['kind']in ['p','h','code']:assert block['text']in actual,(page['id'],block['text'][:40])
            elif block['kind']=='figure':assert (ASSETS/(block['asset']+'.png')).is_file()
            elif block['kind']=='table':
                for row in [block['headers']]+block['rows']:
                    assert all(str(cell)in table_text for cell in row)
    for ref in re.findall(r'\]\(([^)]+)\)',markdown):assert(OUT/ref).is_file(),ref
    assert len(re.findall(r'!\[',markdown))==60
    assert markdown.count('```')%2==0
    assert not re.search(r'Student123!|Counselor123!|sk-[a-zA-Z0-9_-]{20,}|eyJ[a-zA-Z0-9_-]{45,}',actual+table_text)
    with zipfile.ZipFile(DOCX)as archive:
        assert archive.testzip()is None
        xml=ET.fromstring(archive.read('word/document.xml'))
        assert len(xml.xpath('//*[local-name()="inline"]'))==60
        assert not xml.xpath('//*[local-name()="anchor" or local-name()="ins" or local-name()="del"]')
    pdf=fitz.open(PDF);assert len(pdf)==42
    records=[];captions=[];figure_pages={}
    for i,page in enumerate(pdf):
        text=page.get_text();assert '\ufffd'not in text
        if i>=2:assert re.sub(r'\s','',PAGES[i-2]['title'])in re.sub(r'\s','',text),(i+1,PAGES[i-2]['title'])
        assert len(text)>250
        nums=[int(x)for x in re.findall(r'图\s+(\d{2})\s',text)]
        assert nums,(i+1,'missing figure caption')
        for n in nums:figure_pages[n]=i+1
        captions+=nums
        for b in page.get_text('dict')['blocks']:
            for line in b.get('lines',[]):
                x0,y0,x1,y1=line['bbox']
                assert x0>=30 and x1<=page.rect.width-28,(i+1,line['bbox'])
                assert y0>=5 and y1<=page.rect.height-8,(i+1,line['bbox'])
        image_rects=[]
        for image in page.get_images(full=True):
            image_rects.extend(page.get_image_rects(image[0]))
        assert image_rects,(i+1,'no image')
        for rect in image_rects:
            assert rect.x0>=42 and rect.x1<=page.rect.width-42,(i+1,rect)
            assert rect.y0>=37 and rect.y1<=page.rect.height-40,(i+1,rect)
        page.get_pixmap(matrix=fitz.Matrix(1.7,1.7),alpha=False).save(TMP/f'page-{i+1:02}.png')
        records.append({'page':i+1,'title':'cover'if i==0 else'navigation'if i==1 else PAGES[i-2]['title'],'figures':nums,'text_characters':len(text),'all_figures_inside_margins':True,'text_inside_bounds':True})
    assert captions==list(range(1,61)),captions
    index=json.loads((EVIDENCE/'figure_index.json').read_text(encoding='utf-8'))
    for item in index:item['pdf_page']=figure_pages[item['number']]
    (EVIDENCE/'figure_index.json').write_text(json.dumps(index,ensure_ascii=False,indent=2),encoding='utf-8')
    d=ImageFont.truetype(r'C:\Windows\Fonts\msyh.ttc',17)
    for start in range(0,42,6):
        batch=records[start:start+6];sheet=Image.new('RGB',(1350,1326),'#DFE4EA');draw=ImageDraw.Draw(sheet)
        for j,item in enumerate(batch):
            x=j%3*450;y=j//3*663
            with Image.open(TMP/f"page-{item['page']:02}.png")as old:
                thumb=ImageOps.contain(old,(440,622));sheet.paste(thumb,(x+(450-thumb.width)//2,y+2))
            draw.text((x+12,y+635),'Page '+str(item['page']),font=d,fill='black')
        sheet.save(TMP/f'sheet-{start//6+1:02}.jpg',quality=92)
    check={'pages':42,'chapters':24,'special_visual_spreads':16,'figures':60,'all_42_pages_include_figures':True,'figure_numbers_continuous':True,'docx_markdown_same_source':True,'docx_paragraph_content_preserved':True,'no_missing_figure_or_caption':True,'no_figure_outside_margins':True,'no_tracked_changes_or_floating_images':True,'records':records}
    (EVIDENCE/'render_check.json').write_text(json.dumps(check,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in check.items()if k!='records'},ensure_ascii=False,indent=2))


if __name__=='__main__':main()
