"""Build the detailed competition outputs from one shared content source."""
from __future__ import annotations

import hashlib
import json
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

from competition_detailed_content import NAME, DATE, PAGES

PROJECT=Path(r'D:\table\ai_jiuye')
OUTPUT=PROJECT/'docs'/'deliverables'/'比赛提交材料_职航智管_20261025'
ASSETS=OUTPUT/'assets_detailed'
EVIDENCE=OUTPUT/'evidence_detailed'
TMP=PROJECT/'output'/'competition-detailed-20261025'
DOCX=OUTPUT/(NAME+'_作品介绍_详细版.docx')
MD=OUTPUT/(NAME+'_作品介绍_详细版.md')
ROOT_SOURCE=PROJECT/'docs'/'deliverables'/'职航智管_GPT图像PPT提示词包_20261020'
FONT=Path(r'C:\Windows\Fonts\msyh.ttc')
REGULAR='等线'
BLACK='000000'
MUTED='535963'
ACCENT='5B37C8'
REVISION='2026-10-26'
for path in [OUTPUT,ASSETS,EVIDENCE,TMP]:path.mkdir(parents=True,exist_ok=True)
ASSET_MAP={}


def font(run,size=11,bold=False,color=BLACK,name=REGULAR):
    run.font.name=name;run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
    rpr=run._element.get_or_add_rPr()
    rf=rpr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
    for key in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+key),name)
    return run


def para(doc,text='',size=11,bold=False,color=BLACK,after=5,line=17.5,align=None):
    paragraph=doc.add_paragraph()
    fmt=paragraph.paragraph_format
    fmt.space_before=Pt(0);fmt.space_after=Pt(after);fmt.line_spacing=Pt(line)
    fmt.widow_control=False
    if align is not None:paragraph.alignment=align
    for i,line_text in enumerate(text.splitlines() or ['']):
        if i:paragraph.add_run().add_break()
        font(paragraph.add_run(line_text),size,bold,color)
    return paragraph


def grid(doc,headers,rows,widths=None):
    widths=widths or [17.2/len(headers)]*len(headers)
    widths=[w*17.2/sum(widths) for w in widths]
    table=doc.add_table(rows=1,cols=len(headers))
    table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for column,width in zip(table.columns,widths):column.width=Cm(width)
    for i,row in enumerate([headers]+rows):
        cells=table.rows[0].cells if i==0 else table.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,value in enumerate(row):
            cell=cells[j];cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'EEEAF8' if i==0 else ('FFFFFF' if i%2 else 'F8F9FA'));props.append(shading)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                node=OxmlElement('w:'+edge)
                for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:node.set(qn('w:'+k),v)
                borders.append(node)
            props.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge,amount in [('top',70),('bottom',70),('left',95),('right',95)]:
                node=OxmlElement('w:'+edge);node.set(qn('w:w'),str(amount));node.set(qn('w:type'),'dxa');margins.append(node)
            props.append(margins)
            p=cell.paragraphs[0];p.paragraph_format.line_spacing=Pt(14.2);p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.keep_with_next=i==0
            font(p.add_run(str(value)),9.6,i==0)
    para(doc,'',size=1,line=2,after=3)
    return table


def image(doc,asset,caption,width=17,height=7.2):
    path=ASSETS/(asset+'.png')
    with Image.open(path) as im:w,h=im.size
    actual=min(width,height*w/h)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(2);p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(path),width=Cm(actual))
    description=p._p.xpath('.//wp:docPr')
    if description:description[0].set('descr',caption)
    para(doc,caption,size=8.8,color=MUTED,line=12.5,after=6,align=WD_ALIGN_PARAGRAPH.CENTER)


def page_field(paragraph):
    for kind in ['begin','instruction','separate','cache','end']:
        if kind=='instruction':
            node=OxmlElement('w:instrText');node.set(qn('xml:space'),'preserve');node.text=' PAGE '
        elif kind=='cache':node=OxmlElement('w:t');node.text='1'
        else:node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),kind)
        run=OxmlElement('w:r');run.append(node);paragraph._p.append(run)


def diagram(name,nodes,note):
    width,height=2400,430
    canvas=Image.new('RGB',(width,height),'white');draw=ImageDraw.Draw(canvas)
    label_font=ImageFont.truetype(str(FONT),41);note_font=ImageFont.truetype(str(FONT),32)
    margin,gap=32,50
    box_width=int((width-2*margin-gap*(len(nodes)-1))/len(nodes))
    for i,label in enumerate(nodes):
        x=margin+i*(box_width+gap)
        draw.rounded_rectangle((x,50,x+box_width,230),radius=14,fill='#F1EEFA',outline='#72648E',width=3)
        lines=label.split('\n')
        for j,text in enumerate(lines):
            bb=draw.textbbox((0,0),text,font=label_font);y=112+(j-(len(lines)-1)/2)*58
            draw.text((x+(box_width-(bb[2]-bb[0]))/2,y),text,font=label_font,fill='#14141A')
        if i+1<len(nodes):
            a=x+box_width+8;b=a+gap-16
            draw.line((a,140,b,140),fill='#75698B',width=5)
            draw.polygon([(b,140),(b-13,130),(b-13,150)],fill='#75698B')
    bb=draw.textbbox((0,0),note,font=note_font)
    draw.text(((width-(bb[2]-bb[0]))/2,320),note,font=note_font,fill='#535963')
    canvas.save(ASSETS/(name+'.png'))
    ASSET_MAP[name]={'type':'diagram','basis':note,'nodes':nodes}


def prepare_assets():
    shots=ROOT_SOURCE/'截图';raw=PROJECT/'output'/'copyright-20261022'/'screenshots'
    def crop(name,path,box=None):
        with Image.open(path) as im:
            picture=im.convert('RGB')
            if box:picture=picture.crop(box)
            picture.save(ASSETS/(name+'.png'))
        ASSET_MAP[name]={'type':'screenshot','source':str(path.relative_to(PROJECT)),'crop':box,'numbers_or_statuses_changed':False}
    crop('student',raw/'student.png',(45,75,1290,750))
    crop('plan',raw/'plan.png',(45,70,1290,755))
    crop('match',raw/'match.png',(45,70,1290,990))
    crop('agent',shots/'09-agent-assistant.png',(50,60,1630,1010))
    crop('interview_report',shots/'07B-interview-summary.png')
    crop('workplace',shots/'08-workplace.png',(60,170,1580,825))
    crop('intervention',shots/'13A-confirmed-feedback.png')
    crop('employment',raw/'employment.png',(45,70,1290,855))
    for name,file in [('tests','C06-test-results.png'),('tools','C08-tool-inventory.png')]:
        source=ROOT_SOURCE/'图表'/file;shutil.copy2(source,ASSETS/(name+'.png'))
        ASSET_MAP[name]={'type':'chart','source':str(source.relative_to(PROJECT)),'data_scope':'2026-10-20 historical local checks'}
    diagram('journey',['材料与事实','差距与目标','计划与训练','人工确认','结果与复盘'],'每次行动回到原始记录和授权范围；演示样本不代表实际就业成效')
    diagram('evidence_chain',['原始材料','来源片段','结构化分析','动作提案','确认与记录'],'材料、引用、建议和执行结果分别存储，不把生成文本自动当作事实')
    canvas=Image.new('RGB',(2200,900),'white');draw=ImageDraw.Draw(canvas)
    tf=ImageFont.truetype(str(FONT),38);sf=ImageFont.truetype(str(FONT),31)
    boxes=[(80,50,2120,175,'浏览器工作区','学生 / 辅导员 / 受限企业评价入口'),
           (80,265,1060,475,'FastAPI 业务后端','身份与归属 / 状态转换 / 事务与审计'),
           (1140,265,2120,475,'Agent 编排与工具目录','主Agent / 专业Agent / 结构化契约'),
           (80,585,1060,805,'业务数据库与任务记录','画像、任务、申请、消息、工具与事件'),
           (1140,585,2120,805,'按部署模式启用的能力','在线模型；可选 OCR / 向量 / 语音 / 视觉')]
    for x1,y1,x2,y2,title,description in boxes:
        draw.rounded_rectangle((x1,y1,x2,y2),radius=20,fill='#F4F1FA',outline='#82768E',width=3)
        draw.text((x1+28,y1+22),title,font=tf,fill='#17141D')
        draw.text((x1+28,y1+90),description,font=sf,fill='#535963')
    for x,y1,y2 in [(570,176,260),(1630,176,260),(570,476,580),(1630,476,580)]:
        draw.line((x,y1,x,y2),fill='#75698B',width=5);draw.polygon([(x,y2),(x-12,y2-18),(x+12,y2-18)],fill='#75698B')
    draw.line((1061,365,1135,365),fill='#75698B',width=5)
    canvas.save(ASSETS/'architecture.png')
    ASSET_MAP['architecture']={'type':'diagram','basis':['backend/app/main.py','backend/app/tools/registry.py','docker-compose.server-demo.yml']}
    state=Image.new('RGB',(2400,570),'white');draw=ImageDraw.Draw(state)
    labels=['收藏 saved','已投递 applied','笔试 assessment','面试 interview','录用 offered']
    xpositions=[35+i*476 for i in range(5)]
    for x,label in zip(xpositions,labels):
        draw.rounded_rectangle((x,45,x+420,175),radius=12,fill='#F1EEFA',outline='#72648E',width=3)
        bb=draw.textbbox((0,0),label,font=tf)
        draw.text((x+(420-(bb[2]-bb[0]))/2,85),label,font=tf,fill='black')
    for x in xpositions[:-1]:
        draw.line((x+420,110,x+464,110),fill='#75698B',width=4);draw.polygon([(x+464,110),(x+451,100),(x+451,120)],fill='#75698B')
    draw.line((xpositions[1]+210,44,xpositions[1]+210,20,xpositions[3]+210,20,xpositions[3]+210,43),fill='#75698B',width=3)
    draw.text((920,202),'已投递也可直接进入面试',font=sf,fill='#535963')
    for x in xpositions:draw.line((x+210,177,x+210,300),fill='#8B8792',width=2)
    draw.line((245,300,2149,300),fill='#8B8792',width=2)
    draw.line((1200,300,1200,350),fill='#8B8792',width=3)
    draw.rounded_rectangle((860,350,1540,470),radius=12,fill='#F3F4F5',outline='#888888',width=3)
    draw.text((970,385),'关闭 closed 终止流转',font=tf,fill='#242424')
    draw.text((520,510),'每次允许的阶段变更保留事件；非法跳转返回冲突错误',font=sf,fill='#535963')
    state.save(ASSETS/'application_states.png')
    ASSET_MAP['application_states']={'type':'diagram','basis':'backend/app/applications/service.py ALLOWED_TRANSITIONS'}
    (EVIDENCE/'asset_sources.json').write_text(json.dumps(ASSET_MAP,ensure_ascii=False,indent=2),encoding='utf-8')


def create_document():
    doc=Document();section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.left_margin=Cm(1.9);section.right_margin=Cm(1.9)
    section.top_margin=Cm(1.85);section.bottom_margin=Cm(1.75)
    section.header_distance=Cm(.8);section.footer_distance=Cm(.8)
    doc.core_properties.title=NAME+' 作品介绍详细版'
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    for name in ['Normal','Title','Heading 1','Heading 2']:
        if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        style=doc.styles[name];style.font.name=REGULAR;style.font.color.rgb=RGBColor(0,0,0)
        # Word's built-in Title style can carry a blue rule from its default
        # template. Remove it explicitly for a plain document title.
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(11);doc.styles['Title'].font.size=Pt(23)
    header=section.header.paragraphs[0];header.paragraph_format.space_after=Pt(0)
    font(header.add_run('职航智管  比赛作品介绍详细版'),8.4,color=MUTED)
    footer=section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(Cm(17.1),WD_TAB_ALIGNMENT.RIGHT)
    font(footer.add_run(REVISION+'  修订完成  历史测试日期见第17章\t'),8.0,color=MUTED)
    page_field(footer)
    top=para(doc,'技术创新 30%  场景落地 30%  开源治理 20%  长期发展 20%',size=10,color=ACCENT,after=34,line=16)
    top.paragraph_format.space_before=Pt(37)
    title=para(doc,NAME,size=23,bold=True,after=14,line=33);title.style=doc.styles['Title']
    para(doc,'比赛作品介绍详细版',size=18,bold=True,after=22,line=25)
    para(doc,'V1.0  ·  项目架构、业务流程、验证证据与维护方案',size=11,color=MUTED,after=20)
    para(doc,'编制 '+DATE+'  修订 '+REVISION,size=9.4,color=MUTED,after=8,line=14)
    image(doc,'agent','真实界面与独立演示记录，保留原服务状态',width=17,height=9.5)
    para(doc,'阅读说明',size=12,bold=True,after=5)
    para(doc,'本文按照四项评分维度组织材料，提供24章实现、验证与治理说明。正文引用的测试与演示记录保留原日期；本次编制不将历史结果写成新版本的完整验证。',size=11,after=9)
    para(doc,'代码仓库与视频公开链接待提交前填写；评审凭据通过赛事指定安全渠道提供。',size=10.2,color=MUTED)
    for page in PAGES:
        title=para(doc,page['title'],size=17,bold=True,after=10,line=25)
        title.style=doc.styles['Heading 1'];title.paragraph_format.page_break_before=True;title.paragraph_format.keep_with_next=True
        for block in page['blocks']:
            kind=block['kind']
            if kind=='p':para(doc,block['text'])
            elif kind=='h':
                hp=para(doc,block['text'],size=12,bold=True,after=4,line=19)
                hp.style=doc.styles['Heading 2'];hp.paragraph_format.keep_with_next=True
            elif kind=='table':grid(doc,block['headers'],block['rows'],block['widths'])
            elif kind=='figure':image(doc,block['asset'],block['caption'],block['width'],block['height'])
            elif kind=='code':
                paragraph=para(doc,block['text'],size=9.3,color='363947',after=7,line=14)
                for run in paragraph.runs:run.font.name='Consolas';run.font.size=Pt(9.3)
                paragraph.paragraph_format.keep_together=True
        refs=para(doc,'实现或证据依据：'+'；'.join(page['sources']),size=8.1,color=MUTED,after=0,line=11.5)
        refs.paragraph_format.space_before=Pt(5)
    doc.save(DOCX)


def write_markdown():
    lines=['# '+NAME+' 作品介绍详细版','', '版本 V1.0  编制日期 '+DATE+'  修订完成 '+REVISION,'',
    '本文件与详细版Word、PDF使用同一内容源。历史实测与本次核对分别标注；文中的评估、治理和路线内容如尚未落实，均按建议说明。','',
    '## 导读','',*['- '+x['title'] for x in PAGES],'']
    for page in PAGES:
        lines+=['## '+page['title'],'']
        for block in page['blocks']:
            if block['kind']=='p':lines+=[block['text'],'']
            elif block['kind']=='h':lines+=['### '+block['text'],'']
            elif block['kind']=='figure':lines+=[f"![{block['caption']}](assets_detailed/{block['asset']}.png)",'']
            elif block['kind']=='code':lines+=['```text',block['text'],'```','']
            elif block['kind']=='table':
                escaped=lambda text:str(text).replace('|','\\|').replace('\n','<br>')
                lines+=['| '+' | '.join(map(escaped,block['headers']))+' |','| '+' | '.join(['---']*len(block['headers']))+' |']
                lines+=['| '+' | '.join(map(escaped,row))+' |' for row in block['rows']]
                lines+=['']
        lines+=['实现或证据依据：'+'；'.join('`'+p+'`' for p in page['sources']),'']
    MD.write_text('\n'.join(lines),encoding='utf-8')


def evidence():
    paths=sorted({p for page in PAGES for p in page['sources'] if (PROJECT/p).is_file()})
    records=[]
    for path in paths:
        file=PROJECT/path;records.append({'path':path,'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'bytes':file.stat().st_size})
    manifest={'prepared_on':DATE,'revision_completed_on':REVISION,'software':NAME,'content_source':'scripts/competition_detailed_content.py','chapters':len(PAGES),
              'historical_test_date':'2026-10-20','historical_upload_date':'2026-10-22','full_test_suite_rerun_during_document_revision':False,
              'site_check':{'date':DATE,'requests':[{'url':'http://111.228.30.194/','status':200},{'url':'http://111.228.30.194/api/health','status':200,'body':{'status':'ok'}}],'scope':'Unauthenticated read-only home/health requests; no login or online model check'},
              'source_files':records,'body_nonwhitespace_characters':sum(len(re.sub(r'\s','',b.get('text',''))) for p in PAGES for b in p['blocks']),
              'all_content_nonwhitespace_characters':sum(len(re.sub(r'\s','',str(b))) for p in PAGES for b in p['blocks'])}
    (EVIDENCE/'document_basis.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    historical=PROJECT/'output'/'ppt-evidence-2026-10-20'
    for name in ['data_summary.json','backend-tests-current.xml','backend-tests-current.log','frontend-tests.log','api_performance.json','architecture_inventory.json']:
        shutil.copy2(historical/name,EVIDENCE/name)
    raw_regression=json.loads((historical/'browser-tests-serial.json').read_text(encoding='utf-8-sig'))
    summaries=[]
    for suite in raw_regression['suites']:
        for spec in suite.get('specs',[]):
            for test in spec.get('tests',[]):
                message=test.get('results',[{}])[0].get('error',{}).get('message','')
                message=re.sub(r'\x1b\[[0-9;]*[mK]','',message)
                summaries.append({'title':spec['title'],'project':test['projectName'],'status':test['status'],'error_summary':message[:1100]})
    (EVIDENCE/'browser-regression-summary.json').write_text(json.dumps({'date':'2026-10-20','stats':raw_regression['stats'],'tests':summaries,'note':'Source snippets and trace payloads omitted; original evidence remains in the workspace.'},ensure_ascii=False,indent=2),encoding='utf-8')
    prior_generated=EVIDENCE/'browser-tests-serial.json'
    if prior_generated.exists():prior_generated.unlink()


def main():
    prepare_assets();create_document();write_markdown();evidence()
    print(json.dumps({'docx':str(DOCX),'markdown':str(MD),'chapters':len(PAGES),'body_characters':json.loads((EVIDENCE/'document_basis.json').read_text(encoding='utf-8'))['body_nonwhitespace_characters']},ensure_ascii=False))


if __name__=='__main__':main()
